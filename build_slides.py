"""Build a presentation deck for the Premier League salary classification project.

Generates 1-2 slides per section so the presenter can elaborate verbally.
Output: output/salary_classification_slides.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


PRIMARY = RGBColor(0x0B, 0x3D, 0x91)
ACCENT = RGBColor(0x37, 0x00, 0x3C)
LIGHT = RGBColor(0xF2, 0xF4, 0xF8)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)


def add_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent_bar(slide):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.35)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY
    bar.line.fill.background()


def set_text(tf, text, size=18, bold=False, color=DARK):
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = None
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.55), Inches(12.3), Inches(0.8))
    set_text(tb.text_frame, title, size=32, bold=True, color=PRIMARY)
    if subtitle:
        sb = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.25), Inches(12.3), Inches(0.5)
        )
        set_text(sb.text_frame, subtitle, size=16, color=MUTED)


def add_bullets(slide, bullets, left=0.5, top=2.0, width=12.3, height=4.8, size=18):
    tb = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = "•  " + b
        run.font.size = Pt(size)
        run.font.color.rgb = DARK
        run.font.name = "Calibri"
        p.space_after = Pt(8)


def add_footer(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.3))
    set_text(tb.text_frame, text, size=11, color=MUTED)


def add_kpi_row(slide, kpis, top=2.1):
    """Render a row of KPI cards. kpis is list of (label, value)."""
    n = len(kpis)
    total_w = 12.3
    gap = 0.25
    card_w = (total_w - gap * (n - 1)) / n
    left = 0.5
    for label, value in kpis:
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(card_w),
            Inches(1.3),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT
        card.line.color.rgb = PRIMARY
        card.shadow.inherit = False

        v = slide.shapes.add_textbox(
            Inches(left), Inches(top + 0.15), Inches(card_w), Inches(0.6)
        )
        set_text(v.text_frame, value, size=28, bold=True, color=PRIMARY)
        v.text_frame.paragraphs[0].alignment = 2  # center

        l = slide.shapes.add_textbox(
            Inches(left), Inches(top + 0.78), Inches(card_w), Inches(0.45)
        )
        set_text(l.text_frame, label, size=13, color=MUTED)
        l.text_frame.paragraphs[0].alignment = 2

        left += card_w + gap


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_background(slide, RGBColor(0xFF, 0xFF, 0xFF))
    add_accent_bar(slide)
    return slide


def title_slide(prs):
    slide = new_slide(prs)
    add_background(slide, RGBColor(0xFF, 0xFF, 0xFF))
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(2.2), Inches(13.333), Inches(2.6)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = PRIMARY
    band.line.fill.background()

    t = slide.shapes.add_textbox(Inches(0.6), Inches(2.6), Inches(12), Inches(1.2))
    set_text(
        t.text_frame,
        "Premier League Salary Tier Classification",
        size=40,
        bold=True,
        color=RGBColor(0xFF, 0xFF, 0xFF),
    )
    s = slide.shapes.add_textbox(Inches(0.6), Inches(3.6), Inches(12), Inches(0.8))
    set_text(
        s.text_frame,
        "Predicting Low / Medium / High wage tiers from a single FBRef stats season (2024–25)",
        size=20,
        color=RGBColor(0xE6, 0xEC, 0xF8),
    )
    a = slide.shapes.add_textbox(Inches(0.6), Inches(5.4), Inches(12), Inches(0.5))
    set_text(a.text_frame, "Yash  ·  Salary Classification Pipeline", size=16, color=MUTED)
    return slide


def section_divider(prs, number, title, summary):
    slide = new_slide(prs)
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(2.6), Inches(13.333), Inches(2.0)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT
    band.line.fill.background()

    n = slide.shapes.add_textbox(Inches(0.6), Inches(2.75), Inches(12), Inches(0.6))
    set_text(n.text_frame, f"Section {number}", size=18, bold=True, color=RGBColor(0xFF, 0xC8, 0x70))

    t = slide.shapes.add_textbox(Inches(0.6), Inches(3.2), Inches(12), Inches(1.0))
    set_text(t.text_frame, title, size=36, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    s = slide.shapes.add_textbox(Inches(0.6), Inches(4.8), Inches(12), Inches(1.0))
    set_text(s.text_frame, summary, size=18, color=MUTED)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ----- Title -----
    title_slide(prs)

    # ============================================================
    # SECTION 1: Merge Stats and Salaries
    # ============================================================
    section_divider(
        prs,
        1,
        "Merging Stats and Salaries",
        "Joining FBRef performance data with wage data — exact first, fuzzy second.",
    )

    # Slide 1.1 — strategy
    s = new_slide(prs)
    add_title(
        s,
        "Merge strategy: exact first, then fuzzy",
        "Two-pass join keyed on a normalized player name",
    )
    add_bullets(
        s,
        [
            "Normalize names (lowercase, strip accents, drop punctuation) on both sides.",
            "Pass 1 — exact merge on the normalized name.",
            "Pass 2 — only for the leftover unmatched players, run rapidfuzz.token_sort_ratio.",
            "Accept a fuzzy match when the score clears the cutoff defined in CONFIG.",
            "Fill the wage fields from the matched salary row; drop anyone still unmatched (no label = no supervised signal).",
        ],
    )
    add_footer(s, "Section 1 · Data merge")

    # Slide 1.2 — merge stats KPI
    s = new_slide(prs)
    add_title(
        s,
        "Merge results",
        "How many players survived the join, and how (numbers come from the notebook run).",
    )
    add_kpi_row(
        s,
        [
            ("Total players (stats)", "≈ 500"),
            ("Exact matches", "Majority"),
            ("Fuzzy matches", "Small tail"),
            ("Unmatched → dropped", "Few"),
        ],
        top=2.1,
    )
    add_bullets(
        s,
        [
            "Print-out in notebook covers: total players, exact count, fuzzy count, and the unmatched list.",
            "Final modelling set after merge + GK exclusion: 359 players → 287 train / 72 test.",
            "Talking point: dropping unmatched is intentional — they cannot be labelled with a wage tier.",
        ],
        top=3.7,
        height=3.3,
    )
    add_footer(s, "Section 1 · Data merge")

    # ============================================================
    # SECTION 2: Class balance & split
    # ============================================================
    section_divider(
        prs,
        2,
        "Class Balance & Split Strategy",
        "Stratifying by position × tier so train and test see the same mix.",
    )

    # Slide 2.1 — class balance
    s = new_slide(prs)
    add_title(
        s,
        "Class balance by position and tier",
        "Tiers are formed within each position to keep them comparable.",
    )
    add_bullets(
        s,
        [
            "Tiers (Low / Medium / High) are computed via within-position wage percentiles.",
            "Plot of class counts grouped by position × tier exposes any thin cells.",
            "Defenders and midfielders dominate; forwards are smaller but well-paid on average.",
            "This view tells us which combinations risk being under-sampled in CV.",
        ],
    )
    add_footer(s, "Section 2 · Splitting")

    # Slide 2.2 — split strategy
    s = new_slide(prs)
    add_title(
        s,
        "Train / test split (80 / 20)",
        "Stratified on a combined position + tier key.",
    )
    add_kpi_row(
        s,
        [
            ("Train", "287"),
            ("Test", "72"),
            ("Split", "80 / 20"),
            ("Stratify key", "pos_+_tier"),
        ],
        top=2.1,
    )
    add_bullets(
        s,
        [
            "stratify = pos + \"_\" + salary_tier preserves the tier mix inside every position.",
            "Without it, a position with few High-tier players could lose them entirely from train or test.",
            "Same key is reused for the StratifiedKFold splits during cross-validation.",
        ],
        top=3.7,
        height=3.3,
    )
    add_footer(s, "Section 2 · Splitting")

    # ============================================================
    # SECTION 3: Training & selection
    # ============================================================
    section_divider(
        prs,
        3,
        "Training & Model Selection",
        "Three models, shared folds, picked by CV macro-F1.",
    )

    # Slide 3.1 — setup
    s = new_slide(prs)
    add_title(
        s,
        "How models are compared",
        "Shared StratifiedKFold splits + cross_val_score on macro-F1.",
    )
    add_bullets(
        s,
        [
            "Models evaluated: Logistic Regression, Random Forest, XGBoost.",
            "Same StratifiedKFold object passed to all three → identical folds, fair comparison.",
            "Scoring metric: macro-F1 (treats all three tiers equally, important with class imbalance).",
            "Selection rule: pick the model with the highest mean CV macro-F1 on the training set.",
        ],
    )
    add_footer(s, "Section 3 · Modelling")

    # Slide 3.2 — results
    s = new_slide(prs)
    add_title(
        s,
        "Results and selected model",
        "Logistic Regression wins on CV, but the test picture is messier.",
    )

    # Build a results table manually
    headers = ["Model", "CV macro-F1", "Test macro-F1"]
    rows = [
        ["Logistic Regression  (selected)", "0.446", "0.265"],
        ["Random Forest", "0.385", "0.404"],
        ["XGBoost", "0.371", "0.402"],
    ]
    table_shape = s.shapes.add_table(
        len(rows) + 1, len(headers), Inches(0.8), Inches(2.1), Inches(11.7), Inches(2.0)
    )
    tbl = table_shape.table
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        set_text(cell.text_frame, h, size=16, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if i % 2 == 1 else RGBColor(0xFF, 0xFF, 0xFF)
            bold = i == 1
            set_text(cell.text_frame, v, size=15, bold=bold, color=DARK)

    add_bullets(
        s,
        [
            "Selection rule chooses LR on CV mean — but RF/XGB generalise better on the held-out 72.",
            "Talking point: that gap is the classic CV vs test disagreement on small data; we discuss it.",
        ],
        top=4.6,
        height=2.4,
    )
    add_footer(s, "Section 3 · Modelling")

    # ============================================================
    # SECTION 4: Limitations & next steps
    # ============================================================
    section_divider(
        prs,
        4,
        "Limitations & Next Steps",
        "What the dataset cannot tell us — and how to fix it.",
    )

    # Slide 4.1 — limitations
    s = new_slide(prs)
    add_title(
        s,
        "Where the model is weak",
        "Most of the weakness traces back to missing feature families.",
    )
    add_bullets(
        s,
        [
            "Missing feature families: xG, progressive carries/passes, possession, defensive events.",
            "Defenders are likely under-modelled — the available stats reward attacking output.",
            "Position simplification (collapsing into MF/DF/FW) introduces noise across hybrid roles.",
            "Name matching can still drop players even after fuzzy fallback.",
        ],
    )
    add_footer(s, "Section 4 · Limitations")

    # Slide 4.2 — improvements
    s = new_slide(prs)
    add_title(
        s,
        "Suggested improvements",
        "Each addresses one of the limitations above.",
    )
    add_bullets(
        s,
        [
            "Pull richer FBRef tables (shooting, passing, possession, defensive actions).",
            "Constrain fuzzy matching to the same squad → reduces wrong-player collisions.",
            "Treat the target as ordinal (Low < Medium < High) instead of plain multi-class.",
            "Engineer position-specific features so defenders aren’t judged purely on attacking stats.",
        ],
    )
    add_footer(s, "Section 4 · Limitations")

    # ============================================================
    # SECTION 5: Inaccurate classifications
    # ============================================================
    section_divider(
        prs,
        5,
        "Inaccurate Classifications",
        "Where the model gets specific players wrong, and why.",
    )

    s = new_slide(prs)
    add_title(
        s,
        "Anomalies in the predictions",
        "Concrete cases highlight what the feature set can’t see.",
    )
    add_bullets(
        s,
        [
            "Trent Alexander-Arnold flagged as struggling in 2025 — model misses progression / creation stats that justify his wage.",
            "Similar misses happen with deep-lying playmakers and ball-playing centre-backs.",
            "Pattern: when value comes from passing, possession, or chance creation, our features under-rate it.",
            "Adds weight to the next-steps list: richer features and ordinal targets should reduce these specific errors.",
        ],
    )
    add_footer(s, "Section 5 · Anomalies")

    # ============================================================
    # Closing
    # ============================================================
    s = new_slide(prs)
    add_title(s, "Summary", "What we built, what it shows, what’s next.")
    add_bullets(
        s,
        [
            "Pipeline: clean stats → exact + fuzzy salary merge → within-position tiers → stratified split → 3 models.",
            "Selected by CV macro-F1: Logistic Regression (0.446 CV / 0.265 test).",
            "Tree models generalise better on test — a sign that with richer features we should re-evaluate.",
            "Next: add xG / progression / defensive features, ordinal target, squad-aware matching.",
        ],
    )
    add_footer(s, "Thank you · Questions?")

    out_dir = Path("output")
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / "salary_classification_slides.pptx"
    prs.save(out_path)
    print(f"Saved {out_path}  ({out_path.stat().st_size / 1024:.1f} KB)")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
